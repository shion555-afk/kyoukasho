# -*- coding: utf-8 -*-
"""
コース・学科別の購入リスト用 印刷レイアウト共通部品

生徒名簿がない学年（各校の新1年、工業高校の全学年）で使う。
「1コース（学科）＝1ページ」で、教科書一覧・税区分・合計を印刷できる形にする。
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

FONT_NAME = 'Arial'
NORMAL_FONT = Font(name=FONT_NAME, size=10)
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
SCHOOL_FONT = Font(name=FONT_NAME, bold=True, size=12)
COURSE_FONT = Font(name=FONT_NAME, bold=True, size=12)
SMALL_FONT = Font(name=FONT_NAME, size=9)
SMALL_BOLD_FONT = Font(name=FONT_NAME, bold=True, size=9)
TOTAL_LABEL_FONT = Font(name=FONT_NAME, bold=True, size=11)
TOTAL_VALUE_FONT = Font(name=FONT_NAME, bold=True, size=14)

THIN = Side(style='thin', color='999999')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')
CENTER = Alignment(horizontal='center', vertical='center')
RIGHT = Alignment(horizontal='right', vertical='center')

TAX_RATE = 1.1


def is_blank_code(code):
    """教科書番号が空欄・ハイフンなら準教科書（課税）とみなす。"""
    s = str(code).strip() if code is not None else ''
    return s in ('', '-', '－', 'None')


def calc_tax(price, taxable):
    """税込価格から内税分を逆算する。非課税なら0。"""
    if not taxable:
        return 0
    return round(price - price / TAX_RATE)


def make_book(code, publisher, title, price, taxable=None):
    """印刷レイアウト用の1冊分のデータを作る。
    taxable を省略した場合は、教科書番号の有無から自動判定する。"""
    price = int(price) if price else 0
    if taxable is None:
        taxable = is_blank_code(code)
    return {'code': code, 'publisher': publisher, 'title': title,
            'price': price, 'taxable': taxable}


def _write_table(ws, row, books):
    """教科書の表を書き、(次の行, 非課税合計, 課税合計) を返す。"""
    headers = ['教科書番号', '発行所', '書名', '税区分', '税額', '価格']
    for i, h in enumerate(headers):
        cell = ws.cell(row=row, column=2 + i, value=h)
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
        cell.fill = HEADER_FILL
    row += 1

    hikazei = kazei = 0
    for b in books:
        tax = calc_tax(b['price'], b['taxable'])
        if b['taxable']:
            kazei += b['price']
        else:
            hikazei += b['price']
        values = [b['code'], b['publisher'], b['title'],
                  '課' if b['taxable'] else '非', tax, b['price']]
        for i, v in enumerate(values):
            cell = ws.cell(row=row, column=2 + i, value=v)
            cell.font = NORMAL_FONT
            cell.border = BORDER
            if i in (4, 5):
                cell.alignment = RIGHT
            elif i == 3:
                cell.alignment = CENTER
        row += 1
    return row, hikazei, kazei


def build_course_list_print(pages, out_path, doc_title, school_name,
                            sale_date='', notes=None, subtitle=''):
    """コース・学科ごとに1ページの購入リストを作る。

    pages: 以下の形の辞書のリスト
      {
        'course': 'コース名（見出しに出す。空でも可）',
        'books' : [make_book(...), ...],          # 合計に含める
        'choices': [make_book(...), ...],         # いずれか1冊を選ぶ必須科目（合計に加算して併記）
        'choices_label': '芸術（いずれか1冊を選択）',
        'extras': [make_book(...), ...],          # 参考掲載（合計に含めない）
        'extras_label': '（参考）選択科目（芸術）※合計には含めません',
      }
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '購入リスト（印刷用）'

    col_widths = {'A': 3, 'B': 14, 'C': 16, 'D': 46, 'E': 8, 'F': 9, 'G': 10}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    row = 1
    for page in pages:
        ws.cell(row=row, column=2, value=doc_title).font = TITLE_FONT
        row += 1
        if subtitle:
            ws.cell(row=row, column=2, value=subtitle).font = SMALL_FONT
        row += 1

        ws.cell(row=row, column=2, value=school_name).font = SCHOOL_FONT
        if page.get('course'):
            ws.cell(row=row, column=5, value=page['course']).font = COURSE_FONT
        row += 1
        if sale_date:
            ws.cell(row=row, column=2, value=sale_date).font = SMALL_FONT
        row += 1

        books = page.get('books', [])
        base_total = sum(b['price'] for b in books)
        choices = page.get('choices', [])

        ws.cell(row=row, column=2, value='合計金額').font = TOTAL_LABEL_FONT
        if choices:
            lo = base_total + min(c['price'] for c in choices)
            hi = base_total + max(c['price'] for c in choices)
            ws.cell(row=row, column=3,
                    value=(f'{lo:,} 〜 {hi:,} 円' if lo != hi else f'{lo:,} 円')).font = TOTAL_VALUE_FONT
            ws.cell(row=row, column=5,
                    value='※選ぶ科目により金額が変わります').font = SMALL_FONT
        else:
            ws.cell(row=row, column=3, value=base_total).font = TOTAL_VALUE_FONT
            ws.cell(row=row, column=3).alignment = RIGHT
            ws.cell(row=row, column=4, value='円').font = TOTAL_LABEL_FONT
        row += 2

        row, hikazei, kazei = _write_table(ws, row, books)

        if choices:
            row += 1
            ws.cell(row=row, column=2,
                    value=page.get('choices_label', '（いずれか1冊を選択・合計に含みます）')).font = SMALL_BOLD_FONT
            row += 1
            row, _, _ = _write_table(ws, row, choices)

        row += 1
        for label, value in [('非課税合計（教科書）', hikazei),
                             ('課税合計（準教科書）', kazei),
                             ('合計', hikazei + kazei)]:
            ws.cell(row=row, column=5, value=label).font = HEADER_FONT
            cell = ws.cell(row=row, column=7, value=value)
            cell.font = HEADER_FONT
            cell.alignment = RIGHT
            row += 1

        if choices:
            for c in choices:
                ws.cell(row=row, column=5, value=f"上記＋{c['title']}").font = HEADER_FONT
                cell = ws.cell(row=row, column=7, value=hikazei + kazei + c['price'])
                cell.font = HEADER_FONT
                cell.alignment = RIGHT
                row += 1

        extras = page.get('extras', [])
        if extras:
            row += 1
            ws.cell(row=row, column=2,
                    value=page.get('extras_label', '（参考）※合計には含めません')).font = SMALL_BOLD_FONT
            row += 1
            row, _, _ = _write_table(ws, row, extras)

        if notes:
            row += 1
            ws.cell(row=row, column=2, value='【お願い・注意事項】').font = SMALL_BOLD_FONT
            row += 1
            for note in notes:
                ws.cell(row=row, column=2, value=f'・{note}').font = SMALL_FONT
                row += 1

        row += 1
        ws.row_breaks.append(openpyxl.worksheet.pagebreak.Break(id=row - 1))

    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(out_path)
    print(f'保存しました: {out_path}')
    print(f'ページ数（コース・学科数）: {len(pages)}')


def ask_print_options(default_path):
    """コマンド版で、印刷用レイアウトの設定を対話的に聞く。
    作らない場合は None を返す。"""
    ans = input('印刷用レイアウトも作りますか？（y/n）: ').strip().lower()
    if ans != 'y':
        return None
    path = input(f'印刷用ファイルの出力ファイル名（既定: {default_path}）: ').strip() or default_path
    sale_date = input('販売日の文言（空欄可）: ').strip()
    print('注意事項を1行ずつ入力してください（入力なしでEnterのみ押すと終了）:')
    notes = []
    while True:
        note = input(f'  注意事項{len(notes) + 1}: ').strip()
        if not note:
            break
        notes.append(note)
    return {'path': path, 'sale_date': sale_date, 'notes': notes}
