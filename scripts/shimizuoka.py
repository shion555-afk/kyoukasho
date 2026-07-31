# -*- coding: utf-8 -*-
"""
北海道室蘭清水丘高等学校 教科書購入票 自動生成ツール

【必要なファイル】学年ごとに別ファイル
  - 新1年：2026_教科書購入票_新１年_書店提出_ver_2.xlsx
  - 新2年：2026_教科書購入票_新２年_書店提出_ver_2.xlsm
  - 新3年：2026_教科書購入票_新３年_書店提出_ver_2.xlsm

【使い方】
  python shimizuoka.py
"""
import unicodedata
import openpyxl
import printlayout
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

FONT_NAME = 'Arial'
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
NORMAL_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style='thin', color='999999')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')


def style_header_row(ws, row, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER


def price_tax(m, textbook_mark):
    price = int(m['price'])
    if m['is_textbook'] == textbook_mark:
        return price, 0
    tax = round(price - price / 1.1)
    return price, tax


# ============================================================
# 新2年生
# ============================================================

def load_2nen(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True, keep_vba=True)
    ws_p = wb['元購入票']
    master = []
    for r in range(14, 62):
        title = ws_p.cell(row=r, column=6).value
        if not title:
            continue
        master.append({
            'row': r, 'kyoka': ws_p.cell(row=r, column=2).value, 'publisher': ws_p.cell(row=r, column=3).value,
            'code': ws_p.cell(row=r, column=4).value, 'subject': ws_p.cell(row=r, column=5).value,
            'title': title, 'is_textbook': ws_p.cell(row=r, column=11).value, 'price': ws_p.cell(row=r, column=12).value,
        })
    ws_s = wb['教科選択']
    students = []
    for r in range(3, 169):
        sid = ws_s.cell(row=r, column=2).value
        name = ws_s.cell(row=r, column=3).value
        bunkei = ws_s.cell(row=r, column=6).value
        if sid is None or name is None or not isinstance(name, str):
            continue
        flags = [ws_s.cell(row=r, column=8 + i).value for i in range(len(master))]
        students.append({'ID': sid, '氏名': name, '系列': bunkei, 'flags': flags})
    return master, students, '＊'


# ============================================================
# 新3年生
# ============================================================

def load_3nen(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True, keep_vba=True)
    ws_price = wb['全体価格表']
    master = []
    for r in range(2, 33):
        title = ws_price.cell(row=r, column=6).value
        if not title:
            continue
        master.append({
            'row': r, 'kyoka': ws_price.cell(row=r, column=2).value, 'publisher': ws_price.cell(row=r, column=3).value,
            'code': ws_price.cell(row=r, column=4).value, 'subject': ws_price.cell(row=r, column=5).value,
            'title': title, 'is_textbook': ws_price.cell(row=r, column=7).value, 'price': ws_price.cell(row=r, column=8).value,
        })
    ws_s = wb['教科選択']
    students = []
    for r in range(3, 166):
        sid = ws_s.cell(row=r, column=2).value
        name = ws_s.cell(row=r, column=3).value
        bunkei = ws_s.cell(row=r, column=6).value
        if sid is None or name is None or not isinstance(name, str):
            continue
        flags = [ws_s.cell(row=r, column=8 + i).value for i in range(len(master))]
        students.append({'ID': sid, '氏名': name, '系列': bunkei, 'flags': flags})
    return master, students, '教科書'


# ============================================================
# 新1年生
# ============================================================

def load_1nen(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ws = wb['購入票']
    master = []
    for r in range(9, 21):
        master.append({'is_textbook': '教科書', 'publisher': ws.cell(row=r, column=3).value,
                        'code': ws.cell(row=r, column=4).value, 'subject': ws.cell(row=r, column=5).value,
                        'title': ws.cell(row=r, column=6).value, 'price': ws.cell(row=r, column=9).value})
    for r in range(23, 41):
        master.append({'is_textbook': '教材', 'publisher': ws.cell(row=r, column=3).value,
                        'code': None, 'subject': ws.cell(row=r, column=5).value,
                        'title': ws.cell(row=r, column=6).value, 'price': ws.cell(row=r, column=9).value})
    geijutsu = []
    for r in [45, 46]:
        geijutsu.append({'publisher': ws.cell(row=r, column=3).value, 'code': ws.cell(row=r, column=4).value,
                          'subject': ws.cell(row=r, column=5).value, 'title': ws.cell(row=r, column=6).value,
                          'price': ws.cell(row=r, column=9).value})
    return master, geijutsu


# ============================================================
# 出力
# ============================================================

def build_workbook(master, students, textbook_mark, out_path):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = '名簿'
    ws1.append(['ID', '氏名', '系列'])
    style_header_row(ws1, 1, 3)
    for s in students:
        ws1.append([s['ID'], s['氏名'], s['系列']])
        for c in range(1, 4):
            ws1.cell(row=ws1.max_row, column=c).font = NORMAL_FONT
    ws1.column_dimensions['B'].width = 16

    ws2 = wb.create_sheet('集計')
    fixed_headers = ['教科', '科目', '出版社', '教科書番号', '書名', '税区分', '価格']
    ws2.append(fixed_headers + [f"{s['ID']} {s['氏名']}" for s in students])
    style_header_row(ws2, 1, len(fixed_headers) + len(students))
    ws2.freeze_panes = ws2.cell(row=2, column=len(fixed_headers) + 1)
    for i, m in enumerate(master):
        row = [m['kyoka'], m['subject'], m['publisher'], m['code'], m['title'],
               '非' if m['is_textbook'] == textbook_mark else '課', m['price']]
        for s in students:
            row.append('○' if s['flags'][i] == 1 else '')
        ws2.append(row)
        r = ws2.max_row
        for c in range(1, len(fixed_headers) + 1):
            ws2.cell(row=r, column=c).font = NORMAL_FONT
    ws2.column_dimensions['E'].width = 32

    ws3 = wb.create_sheet('購入明細')
    ws3.append(['ID', '氏名', '系列', '区分', '教科書番号', '発行所', '書名', '税区分', '税額', '価格'])
    style_header_row(ws3, 1, 10)
    for s in students:
        first = True
        hikazei_total = 0
        kazei_total = 0
        for i, m in enumerate(master):
            if s['flags'][i] != 1:
                continue
            price, tax = price_tax(m, textbook_mark)
            if m['is_textbook'] == textbook_mark:
                hikazei_total += price
            else:
                kazei_total += price
            ws3.append([s['ID'] if first else None, s['氏名'] if first else None, s['系列'] if first else None,
                        m['subject'], m['code'], m['publisher'], m['title'],
                        '非' if m['is_textbook'] == textbook_mark else '課', tax, price])
            r = ws3.max_row
            for c in range(1, 11):
                ws3.cell(row=r, column=c).font = NORMAL_FONT
            first = False
        ws3.append([None, None, None, None, None, None, None, None, '非課税合計（教科書）', hikazei_total])
        ws3.append([None, None, None, None, None, None, None, None, '課税合計（教材）', kazei_total])
        ws3.append([None, None, None, None, None, None, None, None, '総合計', hikazei_total + kazei_total])
        for rr in range(ws3.max_row - 2, ws3.max_row + 1):
            ws3.cell(row=rr, column=9).font = HEADER_FONT
            ws3.cell(row=rr, column=10).font = HEADER_FONT
    ws3.column_dimensions['B'].width = 16
    ws3.column_dimensions['G'].width = 32

    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"生徒数: {len(students)}  教科書マスター行数: {len(master)}")


def build_1nen_workbook(master, geijutsu, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '購入リスト'
    ws.append(['区分', '科目', '発行所', '教科書番号', '書名', '価格'])
    style_header_row(ws, 1, 6)
    total = 0
    for m in master:
        price = int(m['price'])
        total += price
        ws.append([m['is_textbook'], m['subject'], m['publisher'], m['code'], m['title'], price])
        r = ws.max_row
        for c in range(1, 7):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    ws.append([None, None, None, None, '合計', total])
    ws.cell(row=ws.max_row, column=5).font = HEADER_FONT
    ws.cell(row=ws.max_row, column=6).font = HEADER_FONT
    ws.append([])
    ws.append(['（参考）選択科目（芸術）※入学後に1冊購入、合計には含めない'])
    ws.cell(row=ws.max_row, column=1).font = HEADER_FONT
    ws.append(['区分', '科目', '発行所', '教科書番号', '書名', '価格'])
    style_header_row(ws, ws.max_row, 6)
    for m in geijutsu:
        ws.append(['選択', m['subject'], m['publisher'], m['code'], m['title'], int(m['price'])])
        r = ws.max_row
        for c in range(1, 7):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    ws.column_dimensions['E'].width = 40
    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"合計: {total}円")


# ============================================================
# メイン
# ============================================================

def build_1nen_print_layout(master, geijutsu, out_path, sale_date='', notes=None):
    """新1年：全生徒共通の1ページ印刷用リスト。教科書＝非課税、教材＝課税。"""
    books = [printlayout.make_book(m['code'], m['publisher'], m['title'], m['price'],
                                   taxable=(m['is_textbook'] != '教科書'))
             for m in master]
    extras = [printlayout.make_book(g['code'], g['publisher'], g['title'], g['price'], taxable=False)
              for g in geijutsu]
    pages = [{'course': '全生徒共通', 'books': books, 'extras': extras,
              'extras_label': '（参考）選択科目（芸術）※入学後に1冊購入、合計には含めません'}]
    printlayout.build_course_list_print(
        pages, out_path, '○教科書・問題集　購入リスト', '北海道室蘭清水丘高等学校１学年',
        sale_date=sale_date, notes=notes)


def build_print_layout(master, students, textbook_mark, out_path, school_name, sale_date='', notes=None):
    """実際の購入票に近い「1生徒1ページ」の印刷用レイアウトを作る。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '購入票（印刷用）'

    TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
    SCHOOL_FONT = Font(name=FONT_NAME, bold=True, size=12)
    NAME_FONT = Font(name=FONT_NAME, bold=True, size=12)
    TOTAL_LABEL_FONT = Font(name=FONT_NAME, bold=True, size=11)
    TOTAL_VALUE_FONT = Font(name=FONT_NAME, bold=True, size=14)
    TABLE_HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
    CENTER = Alignment(horizontal='center', vertical='center')
    RIGHT = Alignment(horizontal='right', vertical='center')

    col_widths = {'A': 3, 'B': 10, 'C': 14, 'D': 10, 'E': 42, 'F': 8, 'G': 8, 'H': 10}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    row = 1
    for s in students:
        checked = [(i, m) for i, m in enumerate(master) if s['flags'][i] == 1]
        total = sum(int(m['price']) for _, m in checked)

        ws.cell(row=row, column=2, value='○教科書・問題集　購入票').font = TITLE_FONT
        row += 2

        ws.cell(row=row, column=2, value=school_name).font = SCHOOL_FONT
        ws.cell(row=row, column=5, value=f"{s['ID']}　{s['氏名']}（{s['系列']}）").font = NAME_FONT
        row += 1
        if sale_date:
            ws.cell(row=row, column=2, value=sale_date).font = Font(name=FONT_NAME, size=9)
        row += 1

        ws.cell(row=row, column=2, value='合計金額').font = TOTAL_LABEL_FONT
        ws.cell(row=row, column=3, value=total).font = TOTAL_VALUE_FONT
        ws.cell(row=row, column=3).alignment = RIGHT
        ws.cell(row=row, column=4, value='円').font = TOTAL_LABEL_FONT
        row += 2

        headers = ['教科', '教科書番号', '発行所', '書名', '税区分', '税額', '価格']
        for i, h in enumerate(headers):
            cell = ws.cell(row=row, column=2 + i, value=h)
            cell.font = TABLE_HEADER_FONT
            cell.alignment = CENTER
            cell.border = BORDER
            cell.fill = HEADER_FILL
        row += 1

        hikazei_total = 0
        kazei_total = 0
        for _, m in checked:
            price, tax = price_tax(m, textbook_mark)
            if m['is_textbook'] == textbook_mark:
                hikazei_total += price
            else:
                kazei_total += price
            values = [m['kyoka'], m['code'], m['publisher'], m['title'],
                      '非' if m['is_textbook'] == textbook_mark else '課', tax, price]
            for i, v in enumerate(values):
                cell = ws.cell(row=row, column=2 + i, value=v)
                cell.font = NORMAL_FONT
                cell.border = BORDER
                if i in (5, 6):
                    cell.alignment = RIGHT
            row += 1

        for label, value in [
            ('非課税合計（教科書）', hikazei_total),
            ('課税合計（教材）', kazei_total),
            ('総合計', hikazei_total + kazei_total),
        ]:
            ws.cell(row=row, column=5, value=label).font = TABLE_HEADER_FONT
            cell = ws.cell(row=row, column=8, value=value)
            cell.font = TABLE_HEADER_FONT
            cell.alignment = RIGHT
            row += 1

        if notes:
            row += 1
            ws.cell(row=row, column=2, value='【お願い・注意事項】').font = Font(name=FONT_NAME, bold=True, size=9)
            row += 1
            for note in notes:
                ws.cell(row=row, column=2, value=f"・{note}").font = Font(name=FONT_NAME, size=9)
                row += 1

        row += 1
        ws.row_breaks.append(openpyxl.worksheet.pagebreak.Break(id=row - 1))

    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"生徒数: {len(students)}")


def ask_and_build_print_layout(master, students, textbook_mark, school_name):
    ans = input("実際の購入票に近い印刷用レイアウトも作りますか？（y/n）: ").strip().lower()
    if ans != 'y':
        return
    print_out_path = input("印刷用ファイルの出力ファイル名: ").strip()
    sale_date = input("販売日の文言（空欄可）: ").strip()
    print("注意事項を1行ずつ入力してください（入力なしでEnterのみ押すと終了）:")
    notes = []
    while True:
        note = input(f"  注意事項{len(notes) + 1}: ").strip()
        if not note:
            break
        notes.append(note)
    build_print_layout(master, students, textbook_mark, print_out_path, school_name,
                        sale_date=sale_date, notes=notes)


def main():
    print("=== 清水丘高校 教科書購入票 自動生成 ===")
    print("1: 新1年生（共通リスト）")
    print("2: 新2年生")
    print("3: 新3年生")
    grade = input("学年を選んでください（1/2/3）: ").strip()
    wb_path = input("Excelファイルのパスを入力: ").strip()
    out_path = input("出力ファイル名: ").strip()

    if grade == '1':
        master, geijutsu = load_1nen(wb_path)
        build_1nen_workbook(master, geijutsu, out_path)
        opts = printlayout.ask_print_options('清水丘高校_新1年_購入リスト（印刷用）.xlsx')
        if opts:
            build_1nen_print_layout(master, geijutsu, opts['path'],
                                    sale_date=opts['sale_date'], notes=opts['notes'])
    elif grade == '2':
        master, students, mark = load_2nen(wb_path)
        build_workbook(master, students, mark, out_path)
        ask_and_build_print_layout(master, students, mark, '北海道室蘭清水丘高等学校２学年')
    elif grade == '3':
        master, students, mark = load_3nen(wb_path)
        build_workbook(master, students, mark, out_path)
        ask_and_build_print_layout(master, students, mark, '北海道室蘭清水丘高等学校３学年')
    else:
        print("1・2・3のいずれかを入力してください。")


if __name__ == '__main__':
    main()
