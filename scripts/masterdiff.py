# -*- coding: utf-8 -*-
"""
教科書マスターの照合（供給所データ vs 手元の購入票マスター）

毎年の「マスター更新」を手作業でやると漏れが出るので、
供給所データと突き合わせて「何が変わったか」を洗い出す。

- 同じ年度どうしで使えば：手元のマスターが正しいかの検算
- 前年マスター × 今年の供給所データで使えば：今年の更新箇所の一覧

供給所データに「科目名」列があれば自動で使う。無ければ科目コードを表示する。
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from shortage import SCHOOL_CODES, key_of, load_later_term, find_settings_file

FONT_NAME = 'Arial'
NORMAL_FONT = Font(name=FONT_NAME, size=10)
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
SECTION_FONT = Font(name=FONT_NAME, bold=True, size=11)
SMALL_FONT = Font(name=FONT_NAME, size=9)
ALERT_FONT = Font(name=FONT_NAME, bold=True, size=10, color='C00000')

THIN = Side(style='thin', color='999999')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')
CHANGE_FILL = PatternFill('solid', fgColor='FCE4E4')
NEW_FILL = PatternFill('solid', fgColor='E4F4E4')
CENTER = Alignment(horizontal='center', vertical='center')
RIGHT = Alignment(horizontal='right', vertical='center')

# 供給所データで「科目名」が入っているかもしれない列名の候補
SUBJECT_COLS = ('科目名', '科目', '教科・科目', '科目名称')


def supply_subject(d):
    """供給所データの1行から科目名を取り出す。無ければ科目コードを返す。"""
    for c in SUBJECT_COLS:
        if d.get(c):
            return str(d[c])
    return f"（コード{d.get('科目コード')}）"


def _pick_supply_sheet(wb):
    """複数シートがある場合は「科目名」列を持つシートを優先する。"""
    best = None
    for ws in wb.worksheets:
        hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        if any(h in SUBJECT_COLS for h in hdr):
            return ws
        if best is None:
            best = ws
    return best


def load_supply_all(path, school_label):
    """供給所データを読む。列は先頭行から自動で拾うので、科目名列が増えても動く。
    シートが複数ある場合は「科目名」列を持つシートを優先して使う。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = _pick_supply_sheet(wb)
    ncol = ws.max_column
    hdr = [ws.cell(row=1, column=c).value for c in range(1, ncol + 1)]
    code = SCHOOL_CODES[school_label]
    rows = []
    for r in range(2, ws.max_row + 1):
        d = dict(zip(hdr, [ws.cell(row=r, column=c).value for c in range(1, ncol + 1)]))
        if d.get('学校名') is None or d.get('学校コード') != code:
            continue
        rows.append(d)
    return rows


def build_master_diff(school_label, supply_rows, mine_rows, out_path, later_keys=None):
    """mine_rows: [{'grade','subject','code','publisher','title','price'}] （教科書のみ）"""
    if later_keys is None:
        later_keys = load_later_term(school_label)
    sup = {}
    for d in supply_rows:
        sup.setdefault(key_of(d['書名']), d)

    mine = {}
    for m in mine_rows:
        mine.setdefault(key_of(m['title']), m)

    same, price_changed, only_sup, only_mine, later = [], [], [], [], []
    for k, m in mine.items():
        if k in later_keys and k not in sup:
            later.append(m)
            continue
        if k in sup:
            d = sup[k]
            if (d['定価'] or 0) == (m['price'] or 0):
                same.append((m, d))
            else:
                price_changed.append((m, d))
        else:
            only_mine.append(m)
    for k, d in sup.items():
        if k not in mine:
            only_sup.append(d)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '教科書マスターの照合'
    for col, w in {'A': 3, 'B': 16, 'C': 10, 'D': 12, 'E': 46, 'F': 11, 'G': 11, 'H': 10}.items():
        ws.column_dimensions[col].width = w

    r = 1
    ws.cell(row=r, column=2, value=f'{school_label}高校　教科書マスターの照合').font = TITLE_FONT
    r += 1
    ws.cell(row=r, column=2,
            value='供給所データと、手元の購入票マスター（教科書のみ）を突き合わせた結果').font = SMALL_FONT
    r += 2

    def table(label, rows, kind, fill=None):
        nonlocal r
        ws.cell(row=r, column=2, value=f'{label}　{len(rows)}点').font = SECTION_FONT
        r += 1
        if kind == 'both':
            heads = ['科目（供給所）', '教番', '発行所', '書名', '手元の定価', '供給所の定価', '学年']
        elif kind == 'sup':
            heads = ['科目（供給所）', '教番', '発行所', '書名', '', '定価', '']
        else:
            heads = ['科目（手元）', '教番', '発行所', '書名', '定価', '', '学年']
        for i, h in enumerate(heads):
            c = ws.cell(row=r, column=2 + i, value=h)
            c.font = HEADER_FONT; c.fill = HEADER_FILL; c.alignment = CENTER; c.border = BORDER
        r += 1
        if not rows:
            ws.cell(row=r, column=2, value='（該当なし）').font = NORMAL_FONT
            r += 2
            return
        for x in rows:
            if kind == 'both':
                m, d = x
                vals = [supply_subject(d), d['教番'], d['発行者略称'], m['title'],
                        m['price'], d['定価'], m.get('grade', '')]
            elif kind == 'sup':
                vals = [supply_subject(x), x['教番'], x['発行者略称'], x['書名'], '', x['定価'], '']
            else:
                vals = [x.get('subject', ''), x.get('code', ''), x.get('publisher', ''),
                        x['title'], x['price'], '', x.get('grade', '')]
            for i, v in enumerate(vals):
                c = ws.cell(row=r, column=2 + i, value=v)
                c.font = ALERT_FONT if (kind == 'both' and fill and i in (4, 5)) else NORMAL_FONT
                c.border = BORDER
                if fill:
                    c.fill = fill
                if i in (4, 5):
                    c.alignment = RIGHT
            r += 1
        r += 1

    table('■ 定価が変わっているもの（要更新）', price_changed, 'both', CHANGE_FILL)
    table('■ 供給所にあって手元に無いもの（新規採用・改訂の可能性）', only_sup, 'sup', NEW_FILL)
    table('■ 手元にあって供給所に無いもの（採用終了・書名の表記違いの可能性）', only_mine, 'mine')
    if later:
        table('■ 後期使用（前期は発注しないため供給所データに出てきません）', later, 'mine')
    table('■ 一致（変更なし）', same, 'both')

    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out_path)

    print(f'保存しました: {out_path}')
    print(f'  定価変更 {len(price_changed)}点／供給所のみ {len(only_sup)}点／'
          f'手元のみ {len(only_mine)}点／一致 {len(same)}点／後期使用 {len(later)}点')
    for m, d in price_changed[:10]:
        print(f"     ★{str(m['title'])[:34]:<34} {m['price']}円 → {d['定価']}円")
    return {'price_changed': price_changed, 'only_sup': only_sup,
            'only_mine': only_mine, 'same': same, 'later': later}


# ============================================================
# 学校ごとに「教科書（非課税）」の行を集める
# ============================================================

def _blank(c):
    return (str(c).strip() if c is not None else '') in ('', '-', '－', 'None')


def mine_saka(purchase_path):
    import saka
    out = []
    for grade, loader in [('2年', saka.load_master_2nen), ('3年', saka.load_master_3nen)]:
        for b in loader(purchase_path):
            if b['zeikubun'] == '非':
                out.append({'grade': grade, 'subject': b['subject'], 'code': b['code'],
                            'publisher': b['publisher'], 'title': b['title'], 'price': b['teika']})
    m1, g1 = saka.load_1nen(purchase_path)
    for b in m1 + g1:
        if not _blank(b['code']):
            out.append({'grade': '1年', 'subject': '', 'code': b['code'],
                        'publisher': b['publisher'], 'title': b['title'], 'price': b['price']})
    return out


def mine_touhou(purchase_path):
    import touhou
    out = []
    for grade, loader in [('2年', touhou.load_master_2nen), ('3年', touhou.load_master_3nen)]:
        for b in loader(purchase_path):
            if b['zei'] == '非':
                out.append({'grade': grade, 'subject': b['subject'], 'code': b['code'],
                            'publisher': b['publisher'], 'title': b['title'],
                            'price': touhou.book_price_tax(b)[0]})
    m1, g1 = touhou.load_1nen(purchase_path)
    for b in m1 + g1:
        if not _blank(b['code']):
            out.append({'grade': '1年', 'subject': '', 'code': b['code'],
                        'publisher': b['publisher'], 'title': b['title'], 'price': b['price']})
    return out


def mine_kaisei(wb_path):
    import kaisei
    out = []
    for grade, ldr in [('2年', kaisei.load_2nen), ('3年', kaisei.load_3nen)]:
        m, _, _ = ldr(wb_path)
        for b in m:
            out.append({'grade': grade, 'subject': '', 'code': b['code'],
                        'publisher': b['publisher'], 'title': b['title'], 'price': b['price']})
    for b in kaisei.load_1nen(wb_path):
        out.append({'grade': '1年', 'subject': '', 'code': b['code'],
                    'publisher': b['publisher'], 'title': b['title'], 'price': b['price']})
    return out


def mine_otani(wb_path):
    import otani
    out = []
    for grade, sheet, start, n, cols in [('2年', '新2年購入表', 12, 25, otani.COLS_2NEN),
                                          ('3年', '新3年購入表', 16, 20, otani.COLS_3NEN)]:
        for b in otani.load_master(wb_path, sheet, start, n, cols):
            if b['is_textbook'] == '*':
                out.append({'grade': grade, 'subject': b['subject'], 'code': b['code'],
                            'publisher': b['publisher'], 'title': b['title'], 'price': b['price']})
    m1, art = otani.load_1nen(wb_path)
    for b in m1 + art:
        if not _blank(b['code']):
            out.append({'grade': '1年', 'subject': '', 'code': b['code'],
                        'publisher': b['publisher'], 'title': b['title'], 'price': b['price']})
    return out


def mine_shimizuoka(path2, path3, path1=None):
    import shimizuoka
    out = []
    if path2:
        m, _, _ = shimizuoka.load_2nen(path2)
        for b in m:
            if b['is_textbook'] == '＊':
                out.append({'grade': '2年', 'subject': b['subject'], 'code': b['code'],
                            'publisher': b['publisher'], 'title': b['title'], 'price': b['price']})
    if path3:
        m, _, _ = shimizuoka.load_3nen(path3)
        for b in m:
            if b['is_textbook'] == '教科書':
                out.append({'grade': '3年', 'subject': b['subject'], 'code': b['code'],
                            'publisher': b['publisher'], 'title': b['title'], 'price': b['price']})
    if path1:
        m1, g1 = shimizuoka.load_1nen(path1)
        for b in m1 + g1:
            if b.get('is_textbook') == '教科書' or not _blank(b.get('code')):
                out.append({'grade': '1年', 'subject': b.get('subject', ''), 'code': b.get('code'),
                            'publisher': b['publisher'], 'title': b['title'], 'price': b['price']})
    return out


def mine_kogyo(wb_path):
    import kogyo
    out = []
    for grade, loader in [('1年', kogyo.load_1nen), ('2年', kogyo.load_2nen), ('3年', kogyo.load_3nen)]:
        for b in loader(wb_path):
            if not _blank(b['code']):
                out.append({'grade': grade, 'subject': '', 'code': b['code'],
                            'publisher': b['publisher'], 'title': b['title'], 'price': b['price']})
    return out
