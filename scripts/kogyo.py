# -*- coding: utf-8 -*-
"""
工業高校 教科書購入票 自動生成ツール（学科別・名簿なし）

【必要なファイル】
  - 工業高校令和8年ver2.xlsx（工業１年・工業２年・工業３年シートを含む）

【使い方】
  python kogyo.py
"""
import openpyxl
import printlayout
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

FONT_NAME = 'Arial'
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
NORMAL_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style='thin', color='999999')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')


def style_header_row(ws, row, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER


def build_combined_sheet(wb, master, columns):
    ws = wb.active
    ws.title = '購入リスト'
    headers = ['教科書番号', '発行所', '書名'] + [c[0] for c in columns] + ['価格']
    ws.append(headers)
    style_header_row(ws, 1, len(headers))
    totals = [0] * len(columns)
    for m in master:
        flags = [('○' if matcher(m) else '') for _, matcher in columns]
        if not any(flags):
            continue
        price = int(m['price'])
        for i, f in enumerate(flags):
            if f:
                totals[i] += price
        row = [m['code'], m['publisher'], m['title']] + flags + [price]
        ws.append(row)
        r = ws.max_row
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).font = NORMAL_FONT
            if 4 <= c < 4 + len(columns):
                ws.cell(row=r, column=c).alignment = Alignment(horizontal='center')
    ws.append([None, None, '合計'] + totals + [None])
    for c in range(3, 4 + len(columns)):
        ws.cell(row=ws.max_row, column=c).font = HEADER_FONT
    ws.column_dimensions['C'].width = 40
    return dict(zip([c[0] for c in columns], totals))


def build_kogyo_print_layout(master, columns, out_path, grade_label, sale_date='', notes=None):
    """学科（コース）ごとに1ページの印刷用リストを作る。"""
    pages = []
    for name, matcher in columns:
        books = [printlayout.make_book(m['code'], m['publisher'], m['title'], m['price'])
                 for m in master if matcher(m)]
        pages.append({'course': name, 'books': books})
    zen = {'1': '１', '2': '２', '3': '３'}.get(str(grade_label), str(grade_label))
    printlayout.build_course_list_print(
        pages, out_path, '令和８年度教科書及び準教科書購入リスト',
        f'北海道室蘭工業高等学校{zen}学年',
        sale_date=sale_date, notes=notes)


# ============================================================
# 新1年生：電気科・建設科・電子機械科（●印含む）
# ============================================================

def load_1nen(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ws = wb['工業１年']
    master = []
    for r in range(11, 40):
        title = ws.cell(row=r, column=5).value
        if not title:
            continue
        master.append({
            'row': r, 'publisher': ws.cell(row=r, column=3).value, 'code': ws.cell(row=r, column=4).value,
            'title': title, 'denki': ws.cell(row=r, column=10).value, 'kensetsu': ws.cell(row=r, column=11).value,
            'kiki': ws.cell(row=r, column=12).value, 'price': ws.cell(row=r, column=13).value,
        })
    return master


def columns_1nen():
    return [
        ('電気科', lambda m: bool(m['denki'])),
        ('建設科', lambda m: bool(m['kensetsu'])),
        ('電子機械科', lambda m: bool(m['kiki'])),
    ]


def build_1nen(wb_path, out_path):
    master = load_1nen(wb_path)
    wb = openpyxl.Workbook()
    columns = columns_1nen()
    totals = build_combined_sheet(wb, master, columns)
    wb.save(out_path)
    print(f"保存しました: {out_path}")
    for k, v in totals.items():
        print(f"{k}: {v}円")
    return master, columns


# ============================================================
# 新2年生：電気科・建設科（土木/建築）・電子機械科
# ============================================================

def load_2nen(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ws = wb['工業２年']
    master = []
    for r in range(11, 37):
        title = ws.cell(row=r, column=5).value
        if not title:
            continue
        master.append({
            'row': r, 'publisher': ws.cell(row=r, column=3).value, 'code': ws.cell(row=r, column=4).value,
            'title': title, 'denki': ws.cell(row=r, column=10).value, 'kensetsu': ws.cell(row=r, column=11).value,
            'kiki': ws.cell(row=r, column=12).value, 'price': ws.cell(row=r, column=13).value,
        })
    return master


def columns_2nen():
    return [
        ('電気科', lambda m: bool(m['denki'])),
        ('建設科（土木）', lambda m: m['kensetsu'] in ('○', '◇', '◆')),
        ('建設科（建築）', lambda m: m['kensetsu'] in ('○', '▽')),
        ('電子機械科', lambda m: bool(m['kiki'])),
    ]


def build_2nen(wb_path, out_path):
    master = load_2nen(wb_path)
    wb = openpyxl.Workbook()
    columns = columns_2nen()
    totals = build_combined_sheet(wb, master, columns)
    wb.save(out_path)
    print(f"保存しました: {out_path}")
    for k, v in totals.items():
        print(f"{k}: {v}円")
    return master, columns


# ============================================================
# 新3年生：環境土木科・建築科・電気科(進学/専門)・電子機械科(進学/専門)
# ============================================================

def load_3nen(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    ws = wb['工業３年']
    master = []
    for r in range(11, 27):
        title = ws.cell(row=r, column=5).value
        if not title:
            continue
        master.append({
            'row': r, 'publisher': ws.cell(row=r, column=3).value, 'code': ws.cell(row=r, column=4).value,
            'title': title, 'kando': ws.cell(row=r, column=10).value, 'denki': ws.cell(row=r, column=11).value,
            'kenchiku': ws.cell(row=r, column=12).value, 'kiki': ws.cell(row=r, column=13).value,
            'price': ws.cell(row=r, column=14).value,
        })
    return master


def columns_3nen():
    return [
        ('環境土木科', lambda m: m['kando'] == '○'),
        ('電気科（進学）', lambda m: m['denki'] in ('○', '●')),
        ('電気科（専門）', lambda m: m['denki'] in ('○', '▲')),
        ('建築科', lambda m: m['kenchiku'] == '○'),
        ('電子機械科（進学）', lambda m: m['kiki'] in ('○', '●')),
        ('電子機械科（専門）', lambda m: m['kiki'] in ('○', '▲')),
    ]


def build_3nen(wb_path, out_path):
    master = load_3nen(wb_path)
    wb = openpyxl.Workbook()
    columns = columns_3nen()
    totals = build_combined_sheet(wb, master, columns)
    wb.save(out_path)
    print(f"保存しました: {out_path}")
    for k, v in totals.items():
        print(f"{k}: {v}円")
    return master, columns


# ============================================================
# メイン
# ============================================================

def main():
    print("=== 工業高校 教科書購入票 自動生成 ===")
    print("1: 新1年生")
    print("2: 新2年生")
    print("3: 新3年生")
    grade = input("学年を選んでください（1/2/3）: ").strip()
    wb_path = input("Excelファイルのパスを入力: ").strip()
    out_path = input("出力ファイル名: ").strip()

    if grade == '1':
        master, columns = build_1nen(wb_path, out_path)
    elif grade == '2':
        master, columns = build_2nen(wb_path, out_path)
    elif grade == '3':
        master, columns = build_3nen(wb_path, out_path)
    else:
        print("1・2・3のいずれかを入力してください。")
        return

    opts = printlayout.ask_print_options(f'工業高校_新{grade}年_購入リスト（印刷用）.xlsx')
    if opts:
        build_kogyo_print_layout(master, columns, opts['path'], grade,
                                 sale_date=opts['sale_date'], notes=opts['notes'])


if __name__ == '__main__':
    main()
