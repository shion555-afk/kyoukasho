# -*- coding: utf-8 -*-
"""
海星学院高校 教科書購入票 自動生成ツール

【必要なファイル】
  - 教科書販売Excel（例：02_R8海星教科書販売_新入生_2_3年生_生徒個人別_ver_1.xlsm）
    このファイル1つで完結します（新2年生・新3年生・1年生分すべて含む）

【使い方】
  python kaisei.py
"""
import openpyxl
import printlayout
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

FONT_NAME = 'Arial'
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
NORMAL_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style='thin', color='999999')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')


def style_header_row(ws, row, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER


def book_text(m):
    return f"{m['title']}（{m['price']}円）"


# ============================================================
# 新2年生
# ============================================================

ELECTIVE_INDEXES_2 = {8, 9}  # 物理基礎・生物基礎


def load_2nen(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True, keep_vba=True)
    ws_p = wb['集計新2年生']
    master = []
    for r in range(14, 24):
        master.append({
            'publisher': ws_p.cell(row=r, column=5).value, 'code': ws_p.cell(row=r, column=6).value,
            'title': ws_p.cell(row=r, column=7).value, 'price': ws_p.cell(row=r, column=9).value,
        })
    ws_s = wb['生徒名簿2026_新2年']
    students = []
    for r in range(2, 78):
        sid = ws_s.cell(row=r, column=3).value
        name = ws_s.cell(row=r, column=4).value
        if sid is None or name is None:
            continue
        flags = [ws_s.cell(row=r, column=5 + i).value for i in range(10)]
        sentaku = ws_s.cell(row=r, column=19).value
        students.append({'ID': sid, '氏名': name, '選択科目': sentaku, 'flags': flags})
    return master, students, ELECTIVE_INDEXES_2


# ============================================================
# 新3年生
# ============================================================

ELECTIVE_INDEXES_3 = {2, 3, 4, 5, 6}  # 古典探究・地理探究・世界史探究・化学・生物


def load_3nen(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True, keep_vba=True)
    ws_p = wb['集計新3年生']
    master = []
    for r in range(14, 21):
        master.append({
            'publisher': ws_p.cell(row=r, column=5).value, 'code': ws_p.cell(row=r, column=6).value,
            'title': ws_p.cell(row=r, column=7).value, 'price': ws_p.cell(row=r, column=9).value,
        })
    ws_s = wb['生徒名簿2026_新3年']
    students = []
    for r in range(2, 43):
        sid = ws_s.cell(row=r, column=3).value
        name = ws_s.cell(row=r, column=4).value
        if sid is None or name is None:
            continue
        flags = [ws_s.cell(row=r, column=5 + i).value for i in range(7)]
        sentaku = ' / '.join(
            str(ws_s.cell(row=r, column=15 + i).value) for i in range(3)
            if ws_s.cell(row=r, column=15 + i).value)
        students.append({'ID': sid, '氏名': name, '選択科目': sentaku, 'flags': flags})
    return master, students, ELECTIVE_INDEXES_3


# ============================================================
# 新1年生
# ============================================================

def load_1nen(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True, keep_vba=True)
    ws = wb['集計1年']
    master = []
    for r in range(14, 27):
        title = ws.cell(row=r, column=6).value
        if not title:
            continue
        master.append({
            'publisher': ws.cell(row=r, column=4).value, 'code': ws.cell(row=r, column=5).value,
            'title': title, 'price': ws.cell(row=r, column=7).value,
        })
    return master


# ============================================================
# 出力
# ============================================================

def build_workbook(master, students, elective_indexes, out_path):
    n_cols = len(master)
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = '名簿'
    ws1.append(['ID', '氏名', '選択科目'])
    style_header_row(ws1, 1, 3)
    for s in students:
        ws1.append([s['ID'], s['氏名'], s['選択科目']])
        for c in range(1, 4):
            ws1.cell(row=ws1.max_row, column=c).font = NORMAL_FONT
    ws1.column_dimensions['B'].width = 16

    ws2 = wb.create_sheet('集計')
    fixed_headers = ['発行所', '教科書番号', '書名', '価格']
    ws2.append(fixed_headers + [f"{s['ID']} {s['氏名']}" for s in students])
    style_header_row(ws2, 1, len(fixed_headers) + len(students))
    ws2.freeze_panes = ws2.cell(row=2, column=len(fixed_headers) + 1)
    for i, m in enumerate(master):
        row = [m['publisher'], m['code'], m['title'], m['price']]
        for s in students:
            row.append('○' if s['flags'][i] == 1 else '')
        ws2.append(row)
        r = ws2.max_row
        for c in range(1, len(fixed_headers) + 1):
            ws2.cell(row=r, column=c).font = NORMAL_FONT
    ws2.column_dimensions['C'].width = 32

    ws3 = wb.create_sheet('購入明細')
    ws3.append(['ID', '氏名', '区分', '教科書番号', '発行所', '書名', '価格'])
    style_header_row(ws3, 1, 7)
    for s in students:
        required_books = [master[i] for i in range(n_cols) if i not in elective_indexes and s['flags'][i] == 1]
        elective_books = [master[i] for i in elective_indexes if s['flags'][i] == 1]
        groups = [('必修', required_books), ('選択科目', elective_books)]
        first = True
        total = 0
        for label, books in groups:
            for m in books:
                price = int(m['price'])
                total += price
                ws3.append([s['ID'] if first else None, s['氏名'] if first else None,
                            label, m['code'], m['publisher'], m['title'], price])
                r = ws3.max_row
                for c in range(1, 8):
                    ws3.cell(row=r, column=c).font = NORMAL_FONT
                first = False
        ws3.append([None, None, None, None, None, '合計（非課税）', total])
        ws3.cell(row=ws3.max_row, column=6).font = HEADER_FONT
        ws3.cell(row=ws3.max_row, column=7).font = HEADER_FONT
    ws3.column_dimensions['B'].width = 16
    ws3.column_dimensions['F'].width = 32

    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"生徒数: {len(students)}  教科書マスター行数: {len(master)}")


def build_1nen_workbook(master, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '購入リスト'
    ws.append(['教科書番号', '発行所', '書名', '価格'])
    style_header_row(ws, 1, 4)
    total = 0
    for m in master:
        price = int(m['price'])
        total += price
        ws.append([m['code'], m['publisher'], m['title'], price])
        r = ws.max_row
        for c in range(1, 5):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    ws.append([None, None, '合計', total])
    ws.cell(row=ws.max_row, column=3).font = HEADER_FONT
    ws.cell(row=ws.max_row, column=4).font = HEADER_FONT
    ws.column_dimensions['C'].width = 40
    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"合計: {total}円")


# ============================================================
# メイン
# ============================================================

def build_1nen_print_layout(master, out_path, sale_date='', notes=None):
    """新1年：全生徒共通の1ページ印刷用リスト（全冊 非課税＝教科書）。"""
    books = [printlayout.make_book(m['code'], m['publisher'], m['title'], m['price'], taxable=False)
             for m in master]
    pages = [{'course': '全生徒共通', 'books': books}]
    printlayout.build_course_list_print(
        pages, out_path, '令和８年度教科書購入リスト', '海星学院高等学校１学年',
        sale_date=sale_date, notes=notes)


def build_print_layout(master, students, elective_indexes, out_path, school_name,
                        sale_date='', notes=None):
    """実際の購入票に近い「1生徒1ページ」の印刷用レイアウトを作る（全冊非課税扱い）。"""
    n_cols = len(master)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '購入票（印刷用）'

    TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
    SCHOOL_FONT = Font(name=FONT_NAME, bold=True, size=12)
    NAME_FONT = Font(name=FONT_NAME, bold=True, size=12)
    TOTAL_LABEL_FONT = Font(name=FONT_NAME, bold=True, size=11)
    TOTAL_VALUE_FONT = Font(name=FONT_NAME, bold=True, size=14)
    TABLE_HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
    CENTER = Alignment(horizontal='center', vertical='center')
    RIGHT = Alignment(horizontal='right', vertical='center')

    col_widths = {'A': 3, 'B': 10, 'C': 14, 'D': 42, 'E': 10}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    row = 1
    for s in students:
        required_books = [master[i] for i in range(n_cols) if i not in elective_indexes and s['flags'][i] == 1]
        elective_books = [master[i] for i in elective_indexes if s['flags'][i] == 1]
        groups = [('必修', required_books), ('選択科目', elective_books)]
        total = sum(int(m['price']) for _, books in groups for m in books)

        ws.cell(row=row, column=2, value='令和８年度教科書購入票').font = TITLE_FONT
        row += 2

        ws.cell(row=row, column=2, value=school_name).font = SCHOOL_FONT
        ws.cell(row=row, column=4, value=f"{s['ID']}　{s['氏名']}").font = NAME_FONT
        row += 1
        if sale_date:
            ws.cell(row=row, column=2, value=sale_date).font = Font(name=FONT_NAME, size=9)
        row += 1

        ws.cell(row=row, column=2, value='合計金額').font = TOTAL_LABEL_FONT
        ws.cell(row=row, column=3, value=total).font = TOTAL_VALUE_FONT
        ws.cell(row=row, column=3).alignment = RIGHT
        ws.cell(row=row, column=4, value='円').font = TOTAL_LABEL_FONT
        row += 2

        headers = ['区分', '教科書番号', '発行所', '書名', '価格']
        for i, h in enumerate(headers):
            cell = ws.cell(row=row, column=2 + i, value=h)
            cell.font = TABLE_HEADER_FONT
            cell.alignment = CENTER
            cell.border = BORDER
            cell.fill = HEADER_FILL
        row += 1

        for label, books in groups:
            for m in books:
                values = [label, m['code'], m['publisher'], m['title'], int(m['price'])]
                for i, v in enumerate(values):
                    cell = ws.cell(row=row, column=2 + i, value=v)
                    cell.font = NORMAL_FONT
                    cell.border = BORDER
                    if i == 4:
                        cell.alignment = RIGHT
                row += 1

        if notes:
            row += 1
            ws.cell(row=row, column=2, value='【お願い・注意事項】').font = Font(name=FONT_NAME, bold=True, size=9)
            row += 1
            for note in notes:
                ws.cell(row=row, column=2, value=f"・{note}").font = Font(name=FONT_NAME, size=9)
                row += 1

        row += 1
        ws.row_breaks.append(openpyxl.worksheet.pagebreak.Break(id=row - 1))

    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"生徒数: {len(students)}")


def ask_and_build_print_layout(master, students, elective_indexes, school_name):
    ans = input("実際の購入票に近い印刷用レイアウトも作りますか？（y/n）: ").strip().lower()
    if ans != 'y':
        return
    print_out_path = input("印刷用ファイルの出力ファイル名: ").strip()
    sale_date = input("販売日の文言（空欄可）: ").strip()
    print("注意事項を1行ずつ入力してください（入力なしでEnterのみ押すと終了）:")
    notes = []
    while True:
        note = input(f"  注意事項{len(notes) + 1}: ").strip()
        if not note:
            break
        notes.append(note)
    build_print_layout(master, students, elective_indexes, print_out_path, school_name,
                        sale_date=sale_date, notes=notes)


def main():
    print("=== 海星学院高校 教科書購入票 自動生成 ===")
    print("1: 新2年生")
    print("2: 新3年生")
    print("3: 新1年生（共通リスト）")
    grade = input("学年を選んでください（1/2/3）: ").strip()
    wb_path = input("教科書販売Excelファイルのパスを入力: ").strip()
    out_path = input("出力ファイル名: ").strip()

    if grade == '1':
        master, students, elective = load_2nen(wb_path)
        build_workbook(master, students, elective, out_path)
        ask_and_build_print_layout(master, students, elective, '海星学院高等学校２学年')
    elif grade == '2':
        master, students, elective = load_3nen(wb_path)
        build_workbook(master, students, elective, out_path)
        ask_and_build_print_layout(master, students, elective, '海星学院高等学校３学年')
    elif grade == '3':
        master = load_1nen(wb_path)
        build_1nen_workbook(master, out_path)
        opts = printlayout.ask_print_options('海星学院高校_新1年_購入リスト（印刷用）.xlsx')
        if opts:
            build_1nen_print_layout(master, opts['path'],
                                    sale_date=opts['sale_date'], notes=opts['notes'])
    else:
        print("1・2・3のいずれかを入力してください。")


if __name__ == '__main__':
    main()
