# -*- coding: utf-8 -*-
"""
大谷高校（全日制）教科書購入票 自動生成ツール

【必要なファイル】
  - 教科選択・購入表Excel（例：大谷高校令和8年ver3.xlsm）
    このファイル1つで完結します（新1年・新2年・新3年すべて含む）

【使い方】
  python otani.py
"""
import openpyxl
import printlayout
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

FONT_NAME = 'Arial'
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
NORMAL_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style='thin', color='999999')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')


def style_header_row(ws, row, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER


def book_text(m):
    price = int(m['price'])
    code = m['code'] if m['code'] else '-'
    if m['is_textbook'] == '*':
        return f"[{code}] {m['title']}（{price}円／非課税）"
    tax = round(price - price / 1.1)
    return f"[{code}] {m['title']}（{price}円／税{tax}円）"


# ============================================================
# 新2年生／新3年生（共通処理・列数だけ違う）
# ============================================================

# 購入表シートの列位置。新2年と新3年で1列ずれているので学年ごとに持つ。
COLS_2NEN = {'kyoka': 4, 'publisher': 5, 'code': 6, 'subject': 7,
             'title': 8, 'is_textbook': 10, 'price': 11}
COLS_3NEN = {'kyoka': 3, 'publisher': 4, 'code': 5, 'subject': 6,
             'title': 7, 'is_textbook': 9, 'price': 10}


def load_master(wb_path, sheet_name, start_row, n_rows, cols):
    wb = openpyxl.load_workbook(wb_path, data_only=True, keep_vba=True)
    ws = wb[sheet_name]
    master = []
    for r in range(start_row, start_row + n_rows):
        master.append({
            'row': r,
            'kyoka': ws.cell(row=r, column=cols['kyoka']).value,
            'publisher': ws.cell(row=r, column=cols['publisher']).value,
            'code': ws.cell(row=r, column=cols['code']).value,
            'subject': ws.cell(row=r, column=cols['subject']).value,
            'title': ws.cell(row=r, column=cols['title']).value,
            'is_textbook': ws.cell(row=r, column=cols['is_textbook']).value,
            'price': ws.cell(row=r, column=cols['price']).value,
        })
    return master


def load_students(wb_path, sheet_name, n_cols, start_row=3, end_row_exclusive=None,
                   name_number_fixes=None):
    """出席番号=見本データ(9999)や参照表など、氏名が文字列でない行は自動的に除外します。"""
    wb = openpyxl.load_workbook(wb_path, data_only=True, keep_vba=True)
    ws = wb[sheet_name]
    max_r = end_row_exclusive if end_row_exclusive else ws.max_row + 1
    students = []
    for r in range(start_row, max_r):
        ban = ws.cell(row=r, column=1).value
        sid = ws.cell(row=r, column=2).value
        name = ws.cell(row=r, column=3).value
        ruikei = ws.cell(row=r, column=4).value
        if ban is None or name is None or not isinstance(name, str):
            continue
        flags = [ws.cell(row=r, column=6 + i).value for i in range(n_cols)]
        students.append({'番号': ban, 'ID': sid, '氏名': name, '類型': ruikei, 'flags': flags})
    return students


# ============================================================
# 出力
# ============================================================

def build_workbook(master, students, out_path):
    n_cols = len(master)
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = '名簿'
    ws1.append(['ID', '氏名', '類型'])
    style_header_row(ws1, 1, 3)
    for s in students:
        ws1.append([s['ID'], s['氏名'], s['類型']])
        for c in range(1, 4):
            ws1.cell(row=ws1.max_row, column=c).font = NORMAL_FONT
    ws1.column_dimensions['B'].width = 16

    ws2 = wb.create_sheet('集計')
    fixed_headers = ['教科', '発行所', '教科書番号', '科目', '書名', '価格']
    ws2.append(fixed_headers + [f"{s['ID']} {s['氏名']}" for s in students])
    style_header_row(ws2, 1, len(fixed_headers) + len(students))
    ws2.freeze_panes = ws2.cell(row=2, column=len(fixed_headers) + 1)
    for i, m in enumerate(master):
        row = [m['kyoka'], m['publisher'], m['code'], m['subject'], m['title'], m['price']]
        for s in students:
            row.append('○' if s['flags'][i] == 1 else '')
        ws2.append(row)
        r = ws2.max_row
        for c in range(1, len(fixed_headers) + 1):
            ws2.cell(row=r, column=c).font = NORMAL_FONT
    ws2.column_dimensions['E'].width = 32

    ws3 = wb.create_sheet('購入明細')
    ws3.append(['ID', '氏名', '類型', '区分', '教科書番号', '発行所', '書名', '税区分', '税額', '価格'])
    style_header_row(ws3, 1, 10)
    for s in students:
        first = True
        hikazei_total = 0
        kazei_total = 0
        for i, m in enumerate(master):
            if s['flags'][i] != 1:
                continue
            price = int(m['price'])
            is_hikazei = (m['is_textbook'] == '*')
            tax = 0 if is_hikazei else round(price - price / 1.1)
            if is_hikazei:
                hikazei_total += price
            else:
                kazei_total += price
            ws3.append([s['ID'] if first else None, s['氏名'] if first else None, s['類型'] if first else None,
                        m['subject'], m['code'], m['publisher'], m['title'],
                        '非' if is_hikazei else '課', tax, price])
            r = ws3.max_row
            for c in range(1, 11):
                ws3.cell(row=r, column=c).font = NORMAL_FONT
            first = False
        ws3.append([None, None, None, None, None, None, None, None, '非課税合計（教科書）', hikazei_total])
        ws3.append([None, None, None, None, None, None, None, None, '課税合計（準教科書）', kazei_total])
        ws3.append([None, None, None, None, None, None, None, None, '総合計', hikazei_total + kazei_total])
        for rr in range(ws3.max_row - 2, ws3.max_row + 1):
            ws3.cell(row=rr, column=9).font = HEADER_FONT
            ws3.cell(row=rr, column=10).font = HEADER_FONT
    ws3.column_dimensions['B'].width = 16
    ws3.column_dimensions['G'].width = 32

    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"生徒数: {len(students)}  教科書マスター行数: {len(master)}")


def build_1nen_workbook(master, art, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '購入リスト'
    ws.append(['教科書番号', '発行所', '書名', '難関進路', '文系・理系', '価格'])
    style_header_row(ws, 1, 6)
    nankan_total = 0
    bunkei_total = 0
    for m in master:
        price = int(m['price'])
        if m['nankan'] == '○':
            nankan_total += price
        if m['bunkei_rikei'] == '○':
            bunkei_total += price
        ws.append([m['code'], m['publisher'], m['title'], m['nankan'] or '', m['bunkei_rikei'] or '', price])
        r = ws.max_row
        for c in range(1, 7):
            ws.cell(row=r, column=c).font = NORMAL_FONT
            ws.cell(row=r, column=c).alignment = Alignment(horizontal='center' if c in (4, 5) else 'left')
    ws.append([])
    ws.append(['選択必須（芸術：いずれか1冊を選択、両コース共通・合計に含む）'])
    ws.cell(row=ws.max_row, column=1).font = HEADER_FONT
    for m in art:
        ws.append([m['code'], m['publisher'], m['title'], '選択', '選択', int(m['price'])])
        r = ws.max_row
        for c in range(1, 7):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    ws.append([])
    ws.append([None, None, '難関進路 必修合計', None, None, nankan_total])
    ws.append([None, None, '文系・理系 必修合計', None, None, bunkei_total])
    for a in art:
        price = int(a['price'])
        ws.append([None, None, f"難関進路＋{a['title']}", None, None, nankan_total + price])
        ws.append([None, None, f"文系・理系＋{a['title']}", None, None, bunkei_total + price])
    for rr in range(ws.max_row - 3 - len(art), ws.max_row + 1):
        ws.cell(row=rr, column=3).font = HEADER_FONT
        ws.cell(row=rr, column=6).font = HEADER_FONT
    ws.column_dimensions['C'].width = 32
    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"難関進路必修合計: {nankan_total}円 / 文系・理系必修合計: {bunkei_total}円")


def load_1nen(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True, keep_vba=True)
    ws = wb['新1年']
    master = []
    for r in range(5, 27):
        title = ws.cell(row=r, column=5).value
        if not title:
            continue
        master.append({
            'row': r, 'publisher': ws.cell(row=r, column=3).value, 'code': ws.cell(row=r, column=4).value,
            'title': title, 'nankan': ws.cell(row=r, column=6).value,
            'bunkei_rikei': ws.cell(row=r, column=7).value, 'price': ws.cell(row=r, column=8).value,
        })
    art = []
    for r in [27, 28]:
        art.append({
            'publisher': ws.cell(row=r, column=3).value, 'code': ws.cell(row=r, column=4).value,
            'title': ws.cell(row=r, column=5).value, 'price': ws.cell(row=r, column=8).value,
        })
    return master, art


# ============================================================
# メイン
# ============================================================

def build_1nen_print_layout(master, art, out_path, sale_date='', notes=None):
    """新1年：難関進路・文系理系それぞれ1ページ。芸術は選択必須なので合計に加算して併記。"""
    pages = []
    for label, key in [('難関進路', 'nankan'), ('文系・理系', 'bunkei_rikei')]:
        books = [printlayout.make_book(m['code'], m['publisher'], m['title'], m['price'])
                 for m in master if m[key] == '○']
        choices = [printlayout.make_book(a['code'], a['publisher'], a['title'], a['price'])
                   for a in art]
        pages.append({'course': label, 'books': books, 'choices': choices,
                      'choices_label': '芸術（いずれか1冊を選択・合計に含みます）'})
    printlayout.build_course_list_print(
        pages, out_path, '令和８年度教科書及び準教科書購入リスト', '北海道大谷室蘭高等学校１学年',
        sale_date=sale_date, notes=notes)


def build_print_layout(master, students, out_path, school_name, sale_date='', notes=None):
    """実際の購入票に近い「1生徒1ページ」の印刷用レイアウトを作る。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '購入票（印刷用）'

    TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
    SCHOOL_FONT = Font(name=FONT_NAME, bold=True, size=12)
    NAME_FONT = Font(name=FONT_NAME, bold=True, size=12)
    TOTAL_LABEL_FONT = Font(name=FONT_NAME, bold=True, size=11)
    TOTAL_VALUE_FONT = Font(name=FONT_NAME, bold=True, size=14)
    TABLE_HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
    CENTER = Alignment(horizontal='center', vertical='center')
    RIGHT = Alignment(horizontal='right', vertical='center')

    col_widths = {'A': 3, 'B': 14, 'C': 10, 'D': 14, 'E': 42, 'F': 8, 'G': 8, 'H': 10}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    row = 1
    for s in students:
        checked = [(i, m) for i, m in enumerate(master) if s['flags'][i] == 1]
        total = sum(int(m['price']) for _, m in checked)

        ws.cell(row=row, column=2, value='令和８年度教科書及び準教科書購入表').font = TITLE_FONT
        row += 2

        ws.cell(row=row, column=2, value=school_name).font = SCHOOL_FONT
        ws.cell(row=row, column=5, value=f"{s['ID']}　{s['氏名']}（{s['類型']}）").font = NAME_FONT
        row += 1
        if sale_date:
            ws.cell(row=row, column=2, value=sale_date).font = Font(name=FONT_NAME, size=9)
        row += 1

        ws.cell(row=row, column=2, value='合計金額').font = TOTAL_LABEL_FONT
        ws.cell(row=row, column=3, value=total).font = TOTAL_VALUE_FONT
        ws.cell(row=row, column=3).alignment = RIGHT
        ws.cell(row=row, column=4, value='円').font = TOTAL_LABEL_FONT
        row += 2

        headers = ['教科', '教科書番号', '発行所', '書名', '税区分', '税額', '価格']
        for i, h in enumerate(headers):
            cell = ws.cell(row=row, column=2 + i, value=h)
            cell.font = TABLE_HEADER_FONT
            cell.alignment = CENTER
            cell.border = BORDER
            cell.fill = HEADER_FILL
        row += 1

        hikazei_total = 0
        kazei_total = 0
        for _, m in checked:
            price = int(m['price'])
            is_hikazei = (m['is_textbook'] == '*')
            tax = 0 if is_hikazei else round(price - price / 1.1)
            if is_hikazei:
                hikazei_total += price
            else:
                kazei_total += price
            values = [m['kyoka'], m['code'], m['publisher'], m['title'],
                      '非' if is_hikazei else '課', tax, price]
            for i, v in enumerate(values):
                cell = ws.cell(row=row, column=2 + i, value=v)
                cell.font = NORMAL_FONT
                cell.border = BORDER
                if i in (5, 6):
                    cell.alignment = RIGHT
            row += 1

        for label, value in [
            ('非課税合計（教科書）', hikazei_total),
            ('課税合計（準教科書）', kazei_total),
            ('総合計', hikazei_total + kazei_total),
        ]:
            ws.cell(row=row, column=5, value=label).font = TABLE_HEADER_FONT
            cell = ws.cell(row=row, column=8, value=value)
            cell.font = TABLE_HEADER_FONT
            cell.alignment = RIGHT
            row += 1

        if notes:
            row += 1
            ws.cell(row=row, column=2, value='【お願い・注意事項】').font = Font(name=FONT_NAME, bold=True, size=9)
            row += 1
            for note in notes:
                ws.cell(row=row, column=2, value=f"・{note}").font = Font(name=FONT_NAME, size=9)
                row += 1

        row += 1
        ws.row_breaks.append(openpyxl.worksheet.pagebreak.Break(id=row - 1))

    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"生徒数: {len(students)}")


def ask_and_build_print_layout(students, master, school_name):
    ans = input("実際の購入票に近い印刷用レイアウトも作りますか？（y/n）: ").strip().lower()
    if ans != 'y':
        return
    print_out_path = input("印刷用ファイルの出力ファイル名: ").strip()
    sale_date = input("販売日の文言（空欄可）: ").strip()
    print("注意事項を1行ずつ入力してください（入力なしでEnterのみ押すと終了）:")
    notes = []
    while True:
        note = input(f"  注意事項{len(notes) + 1}: ").strip()
        if not note:
            break
        notes.append(note)
    build_print_layout(master, students, print_out_path, school_name, sale_date=sale_date, notes=notes)


def main():
    print("=== 大谷高校（全日制） 教科書購入票 自動生成 ===")
    print("1: 新1年生（共通リスト）")
    print("2: 新2年生")
    print("3: 新3年生")
    grade = input("学年を選んでください（1/2/3）: ").strip()
    wb_path = input("Excelファイルのパスを入力: ").strip()
    out_path = input("出力ファイル名: ").strip()

    if grade == '1':
        master, art = load_1nen(wb_path)
        build_1nen_workbook(master, art, out_path)
        opts = printlayout.ask_print_options('大谷高校_新1年_購入リスト（印刷用）.xlsx')
        if opts:
            build_1nen_print_layout(master, art, opts['path'],
                                    sale_date=opts['sale_date'], notes=opts['notes'])
    elif grade == '2':
        master = load_master(wb_path, '新2年購入表', 12, 25, COLS_2NEN)
        students = load_students(wb_path, '新2年教科選択', 25, start_row=3)
        build_workbook(master, students, out_path)
        ask_and_build_print_layout(students, master, '北海道大谷室蘭高等学校２学年')
    elif grade == '3':
        master = load_master(wb_path, '新3年購入表', 16, 20, COLS_3NEN)
        students = load_students(wb_path, '新3年教科選択', 20, start_row=3)
        build_workbook(master, students, out_path)
        ask_and_build_print_layout(students, master, '北海道大谷室蘭高等学校３学年')
    else:
        print("1・2・3のいずれかを入力してください。")


if __name__ == '__main__':
    main()
