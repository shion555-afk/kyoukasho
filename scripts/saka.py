# -*- coding: utf-8 -*-
"""
室蘭栄高校 教科書購入票 自動生成ツール

【必要なファイル】
  - 前年（または当年）の購入票ファイル（例：栄高校令和8年ver2.xlsm）
      → 新2年集計・新3年集計・新1年 シートをマスター表として使用
  - 新2年生・新3年生の生データ（例：②新２学年_科目選択一覧_.xlsx など）
      → クラス別シート（1組〜5組）に生徒の選択科目が入っているもの

【使い方】
  python saka.py

コンソールの案内に従って、学年とファイルパスを入力してください。
"""
import re
import unicodedata
import openpyxl
import printlayout
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

FONT_NAME = 'Arial'
HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
NORMAL_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style='thin', color='999999')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')

CHIREKI_SUBJECTS = {'地理探究', '日本史探究', '世界史探究'}


def style_header_row(ws, row, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER


# ============================================================
# 2年生
# ============================================================

def load_master_2nen(purchase_wb_path):
    wb = openpyxl.load_workbook(purchase_wb_path, data_only=True, keep_vba=True)
    ws = wb['新2年集計']
    master = []
    for r in range(3, 55):
        subject = ws.cell(row=r, column=2).value
        if not subject:
            continue
        master.append({
            'row': r, 'subject': subject,
            'publisher': ws.cell(row=r, column=4).value,
            'code': ws.cell(row=r, column=5).value,
            'title': ws.cell(row=r, column=6).value,
            'futsu': ws.cell(row=r, column=7).value,
            'risu': ws.cell(row=r, column=8).value,
            'zeikubun': ws.cell(row=r, column=10).value,
            'honntai': ws.cell(row=r, column=11).value,
            'zei': ws.cell(row=r, column=12).value,
            'teika': ws.cell(row=r, column=13).value,
        })
    return master


def load_students_2nen(raw_wb_path):
    wb = openpyxl.load_workbook(raw_wb_path, data_only=True)
    roster_ws = wb['年度始作業']
    roster = {}
    for row in roster_ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        year, kumi, ban, name = row[0], row[1], row[2], row[3]
        roster[(kumi, ban)] = name

    configs = [
        ('1組', '普通科', [10, 11, 12]), ('2組', '普通科', [10, 11, 12]), ('3組', '普通科', [10, 11, 12]),
        ('4組', '理数科', [8]), ('5組', '理数科', [8]),
    ]
    students = []
    for sheetname, course, cols in configs:
        ws = wb[sheetname]
        kumi_num = int(sheetname[0])
        for r in range(10, ws.max_row + 1):
            ban = ws.cell(row=r, column=2).value
            name_in_sheet = ws.cell(row=r, column=4).value
            if ban is None or name_in_sheet is None:
                continue
            name = roster.get((kumi_num, ban), name_in_sheet)
            subjects = []
            for c in cols:
                v = ws.cell(row=r, column=c).value
                if v not in (None, ''):
                    subjects.append(v)
            if not subjects:
                continue  # 転出・退学者として除外
            students.append({'組': kumi_num, '番': ban, '氏名': name, '系列': course, '科目リスト': subjects})
    students.sort(key=lambda s: (s['組'], s['番']))
    return students


def is_checked_2nen(student, m):
    course = student['系列']
    required = (course == '普通科' and m['futsu'] == '○') or (course == '理数科' and m['risu'] == '○')
    elective = m['subject'] in student['科目リスト']
    if course == '理数科' and m['subject'] in CHIREKI_SUBJECTS:
        elective = False  # 理数科は地歴を前期購入しない
    return required or elective


# ============================================================
# 3年生
# ============================================================

SPECIAL_EXACT_3 = {
    '理数数学Ⅱ(文)': '数学発展',
    '理数生物(理)': '生物・理数生物',
    '理数生物(文)': '生物探究・理数生物',
}
SPECIAL_STRIPPED_3 = {
    '数学ⅢC': '数学Ⅲ',
    '日本史発展': '日本史探究',
}
ROW_OVERRIDE_TITLES_3 = {
    '理数数学Ⅱ(理)': {'PRIME 数学Ⅲ 問題編', 'PRIME 数学Ⅲ 解答編'},
}


def normalize_3(raw):
    if raw in SPECIAL_EXACT_3:
        return SPECIAL_EXACT_3[raw]
    s = re.sub(r'[（(][^）)]*[）)]$', '', raw)
    return SPECIAL_STRIPPED_3.get(s, s)


def subject_matches_3(raw, m, course):
    if raw in ROW_OVERRIDE_TITLES_3:
        return m['title'] in ROW_OVERRIDE_TITLES_3[raw]
    norm = normalize_3(raw)
    master_subject = m['subject']
    if master_subject == '文学探究・文学国語':
        if course == '普通科':
            return norm == '文学探究'
        return norm in ('文学探究', '文学国語')
    if norm == master_subject:
        return True
    if raw in SPECIAL_EXACT_3:
        return False
    return norm in master_subject.split('・')


def load_master_3nen(purchase_wb_path):
    wb = openpyxl.load_workbook(purchase_wb_path, data_only=True, keep_vba=True)
    ws = wb['新3年集計']
    master = []
    r = 3
    while True:
        subject = ws.cell(row=r, column=2).value
        if subject is None or subject == '合計冊数と金額':
            break
        master.append({
            'row': r, 'subject': subject,
            'publisher': ws.cell(row=r, column=4).value,
            'code': ws.cell(row=r, column=5).value,
            'title': ws.cell(row=r, column=6).value,
            'futsu': ws.cell(row=r, column=7).value,
            'risu': ws.cell(row=r, column=8).value,
            'zeikubun': ws.cell(row=r, column=10).value,
            'honntai': ws.cell(row=r, column=11).value,
            'zei': ws.cell(row=r, column=12).value,
            'teika': ws.cell(row=r, column=13).value,
        })
        r += 1
    return master


def load_students_3nen(raw_wb_path):
    wb = openpyxl.load_workbook(raw_wb_path, data_only=True)
    roster_ws = wb['年度始作業']
    roster = {}
    for row in roster_ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        old_kumi, old_ban, new_kumi, new_ban, name = row[0], row[1], row[2], row[3], row[4]
        roster[(new_kumi, new_ban)] = name

    configs = [
        ('1組', '普通科', list(range(26, 33))), ('2組', '普通科', list(range(26, 33))),
        ('3組', '普通科', list(range(26, 33))),
        ('4組', '理数科', list(range(18, 25))), ('5組', '理数科', list(range(18, 25))),
    ]
    students = []
    for sheetname, course, cols in configs:
        ws = wb[sheetname]
        kumi_num = int(sheetname[0])
        for r in range(10, ws.max_row + 1):
            ban = ws.cell(row=r, column=2).value
            name_in_sheet = ws.cell(row=r, column=4).value
            if ban is None or name_in_sheet is None:
                continue
            name = roster.get((kumi_num, ban), name_in_sheet)
            subjects = []
            for c in cols:
                v = ws.cell(row=r, column=c).value
                if v not in (None, ''):
                    subjects.append(v)
            if not subjects:
                continue
            students.append({'組': kumi_num, '番': ban, '氏名': name, '系列': course, '科目リスト': subjects})
    students.sort(key=lambda s: (s['組'], s['番']))
    return students


def is_checked_3nen(student, m):
    course = student['系列']
    required = (course == '普通科' and m['futsu'] == '○') or (course == '理数科' and m['risu'] == '○')
    elective = any(subject_matches_3(raw, m, course) for raw in student['科目リスト'])
    return required or elective


# ============================================================
# 1年生（共通リスト・名簿なし）
# ============================================================

def load_1nen(purchase_wb_path):
    wb = openpyxl.load_workbook(purchase_wb_path, data_only=True, keep_vba=True)
    ws = wb['新1年']
    master = []
    for r in range(9, 65):
        title = ws.cell(row=r, column=5).value
        if not title:
            continue
        master.append({
            'row': r, 'publisher': ws.cell(row=r, column=3).value, 'code': ws.cell(row=r, column=4).value,
            'title': title, 'futsu': ws.cell(row=r, column=6).value, 'risu': ws.cell(row=r, column=7).value,
            'price': ws.cell(row=r, column=8).value,
        })
    geijutsu = []
    for r in [68, 69]:
        geijutsu.append({
            'publisher': ws.cell(row=r, column=3).value, 'code': ws.cell(row=r, column=4).value,
            'title': ws.cell(row=r, column=5).value, 'price': ws.cell(row=r, column=8).value,
        })
    return master, geijutsu


# ============================================================
# 出力（共通処理）
# ============================================================

def build_workbook(master, students, out_path, is_checked_fn):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = '名簿'
    ws1.append(['組', '番', '氏名', '系列'])
    style_header_row(ws1, 1, 4)
    for s in students:
        ws1.append([s['組'], s['番'], s['氏名'], s['系列']])
        for c in range(1, 5):
            ws1.cell(row=ws1.max_row, column=c).font = NORMAL_FONT
    ws1.column_dimensions['C'].width = 16

    ws2 = wb.create_sheet('集計')
    fixed_headers = ['科目', '発行所', '教科書番号', '書名', '本体価格', '税', '定価']
    ws2.append(fixed_headers + [f"{s['組']}組{s['番']}番 {s['氏名']}" for s in students])
    style_header_row(ws2, 1, len(fixed_headers) + len(students))
    ws2.freeze_panes = ws2.cell(row=2, column=len(fixed_headers) + 1)
    for m in master:
        row = [m['subject'], m['publisher'], m['code'], m['title'],
               int(m['honntai']) if m['honntai'] else None, int(m['zei']) if m['zei'] else None,
               int(m['teika']) if m['teika'] else None]
        for s in students:
            row.append('○' if is_checked_fn(s, m) else '')
        ws2.append(row)
        r = ws2.max_row
        for c in range(1, len(fixed_headers) + 1):
            ws2.cell(row=r, column=c).font = NORMAL_FONT
    ws2.column_dimensions['A'].width = 20
    ws2.column_dimensions['D'].width = 32

    ws3 = wb.create_sheet('購入明細')
    ws3.append(['組', '番', '氏名', '系列', '教科書番号', '発行所', '書名', '税区分', '税額', '定価'])
    style_header_row(ws3, 1, 10)
    for s in students:
        hikazei_total = 0
        kazei_total = 0
        first = True
        checked_books = [m for m in master if is_checked_fn(s, m)]
        for m in checked_books:
            teika = int(m['teika']) if m['teika'] else 0
            zeikubun = m['zeikubun']
            zei = int(m['zei']) if m['zei'] else 0
            if zeikubun == '非':
                hikazei_total += teika
            else:
                kazei_total += teika
            ws3.append([
                s['組'] if first else None, s['番'] if first else None, s['氏名'] if first else None,
                s['系列'] if first else None, m['code'], m['publisher'], m['title'], zeikubun, zei, teika,
            ])
            for c in range(1, 11):
                ws3.cell(row=ws3.max_row, column=c).font = NORMAL_FONT
            first = False
        ws3.append([None, None, None, None, None, None, None, None, '非課税合計（教科書）', hikazei_total])
        ws3.append([None, None, None, None, None, None, None, None, '課税合計（準教科書）', kazei_total])
        ws3.append([None, None, None, None, None, None, None, None, '総合計', hikazei_total + kazei_total])
        for rr in range(ws3.max_row - 2, ws3.max_row + 1):
            ws3.cell(row=rr, column=9).font = HEADER_FONT
            ws3.cell(row=rr, column=10).font = HEADER_FONT
    ws3.column_dimensions['C'].width = 16
    ws3.column_dimensions['G'].width = 30

    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"生徒数: {len(students)}  教科書マスター行数: {len(master)}")


def build_1nen_workbook(master, geijutsu, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '購入リスト'
    ws.append(['教科書番号', '発行所', '書名', '普通科', '理数科', '価格'])
    style_header_row(ws, 1, 6)
    futsu_total = 0
    risu_total = 0
    for m in master:
        price = int(m['price'])
        if m['futsu'] == '○':
            futsu_total += price
        if m['risu'] == '○':
            risu_total += price
        ws.append([m['code'], m['publisher'], m['title'], m['futsu'] or '', m['risu'] or '', price])
        r = ws.max_row
        for c in range(1, 7):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    ws.append([None, None, '普通科合計', None, None, futsu_total])
    ws.append([None, None, '理数科合計', None, None, risu_total])
    for rr in (ws.max_row - 1, ws.max_row):
        ws.cell(row=rr, column=3).font = HEADER_FONT
        ws.cell(row=rr, column=6).font = HEADER_FONT
    ws.append([])
    ws.append(['（参考）選択科目（芸術科）※入学後に学校で別売り、合計には含めない'])
    ws.cell(row=ws.max_row, column=1).font = HEADER_FONT
    ws.append(['教科書番号', '発行所', '書名', None, None, '価格'])
    style_header_row(ws, ws.max_row, 6)
    for m in geijutsu:
        ws.append([m['code'], m['publisher'], m['title'], None, None, int(m['price'])])
        r = ws.max_row
        for c in range(1, 7):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    ws.column_dimensions['C'].width = 40
    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"普通科合計: {futsu_total}円 / 理数科合計: {risu_total}円")


# ============================================================
# メイン
# ============================================================

def build_print_layout(master, students, out_path, school_name, is_checked_fn,
                        sale_date='', notes=None):
    """実際の購入票に近い「1生徒1ページ」の印刷用レイアウトを作る。

    sale_date: 販売日の文言（例：'販売日　３月１４日（土）（予備日３月１７日（火））'）
    notes: 注意事項のリスト（例：['クレジットカードでは購入できません。', '紙袋またはカバンをご持参ください。']）
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '購入票（印刷用）'

    TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
    SCHOOL_FONT = Font(name=FONT_NAME, bold=True, size=12)
    NAME_FONT = Font(name=FONT_NAME, bold=True, size=12)
    TOTAL_LABEL_FONT = Font(name=FONT_NAME, bold=True, size=11)
    TOTAL_VALUE_FONT = Font(name=FONT_NAME, bold=True, size=14)
    TABLE_HEADER_FONT = Font(name=FONT_NAME, bold=True, size=10)
    CENTER = Alignment(horizontal='center', vertical='center')
    RIGHT = Alignment(horizontal='right', vertical='center')

    col_widths = {'A': 3, 'B': 14, 'C': 10, 'D': 42, 'E': 10, 'F': 8, 'G': 10}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    row = 1
    for s in students:
        checked_books = [m for m in master if is_checked_fn(s, m)]
        total = sum(int(m['teika']) if m['teika'] else 0 for m in checked_books)

        block_start = row
        ws.cell(row=row, column=2, value='令和８年度教科書及び準教科書購入票').font = TITLE_FONT
        row += 1
        ws.cell(row=row, column=2,
                value='（令和8年1月文部科学大臣が認可し官報で告示した定価）').font = Font(name=FONT_NAME, size=9)
        row += 2

        ws.cell(row=row, column=2, value=school_name).font = SCHOOL_FONT
        ws.cell(row=row, column=5, value=f"{s['組']}組{s['番']}番　{s['氏名']}").font = NAME_FONT
        row += 1
        if sale_date:
            ws.cell(row=row, column=2, value=sale_date).font = Font(name=FONT_NAME, size=9)
        row += 1

        ws.cell(row=row, column=2, value='合計金額').font = TOTAL_LABEL_FONT
        ws.cell(row=row, column=3, value=total).font = TOTAL_VALUE_FONT
        ws.cell(row=row, column=3).alignment = RIGHT
        ws.cell(row=row, column=4, value='円').font = TOTAL_LABEL_FONT
        row += 2

        header_row = row
        headers = ['発行所', '番号', '書名', '金額', '税', '小計']
        for i, h in enumerate(headers):
            cell = ws.cell(row=header_row, column=2 + i, value=h)
            cell.font = TABLE_HEADER_FONT
            cell.alignment = CENTER
            cell.border = BORDER
            cell.fill = HEADER_FILL
        row += 1

        hikazei_total = 0
        kazei_total = 0
        for m in checked_books:
            teika = int(m['teika']) if m['teika'] else 0
            zei = int(m['zei']) if m['zei'] else 0
            honntai = int(m['honntai']) if m['honntai'] else teika
            if m['zeikubun'] == '非':
                hikazei_total += teika
            else:
                kazei_total += teika
            values = [m['publisher'], m['code'], m['title'], honntai, zei, teika]
            for i, v in enumerate(values):
                cell = ws.cell(row=row, column=2 + i, value=v)
                cell.font = NORMAL_FONT
                cell.border = BORDER
                if i in (3, 4, 5):
                    cell.alignment = RIGHT
            row += 1

        for label, value in [
            ('非課税合計（教科書）', hikazei_total),
            ('課税合計（準教科書）', kazei_total),
            ('総合計', hikazei_total + kazei_total),
        ]:
            ws.cell(row=row, column=6, value=label).font = TABLE_HEADER_FONT
            cell = ws.cell(row=row, column=7, value=value)
            cell.font = TABLE_HEADER_FONT
            cell.alignment = RIGHT
            row += 1

        if notes:
            row += 1
            ws.cell(row=row, column=2, value='【お願い・注意事項】').font = Font(name=FONT_NAME, bold=True, size=9)
            row += 1
            for note in notes:
                ws.cell(row=row, column=2, value=f"・{note}").font = Font(name=FONT_NAME, size=9)
                row += 1

        ws.row_dimensions[row].height = 10
        row += 1
        # このページの終わりで改ページ
        ws.row_breaks.append(openpyxl.worksheet.pagebreak.Break(id=row - 1))

    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(out_path)
    print(f"保存しました: {out_path}")
    print(f"生徒数: {len(students)}")


def build_1nen_print_layout(master, geijutsu, out_path, sale_date='', notes=None):
    """新1年：普通科・理数科それぞれ1ページの印刷用リストを作る。"""
    pages = []
    for label, key in [('普通科', 'futsu'), ('理数科', 'risu')]:
        books = [printlayout.make_book(m['code'], m['publisher'], m['title'], m['price'])
                 for m in master if m[key] == '○']
        extras = [printlayout.make_book(g['code'], g['publisher'], g['title'], g['price'])
                  for g in geijutsu]
        pages.append({'course': label, 'books': books, 'extras': extras,
                      'extras_label': '（参考）選択科目（芸術科）※入学後に学校で別売り、合計には含めません'})
    printlayout.build_course_list_print(
        pages, out_path, '令和８年度教科書及び準教科書購入リスト', '北海道室蘭栄高等学校１学年',
        sale_date=sale_date, notes=notes,
        subtitle='（令和8年1月文部科学大臣が認可し官報で告示した定価）')


def ask_and_build_print_layout(master, students, is_checked_fn, school_name):
    ans = input("実際の購入票に近い印刷用レイアウトも作りますか？（y/n）: ").strip().lower()
    if ans != 'y':
        return
    print_out_path = input("印刷用ファイルの出力ファイル名: ").strip()
    sale_date = input("販売日の文言（例: 販売日　３月１４日（土）（予備日３月１７日（火））、空欄可）: ").strip()
    print("注意事項を1行ずつ入力してください（入力なしでEnterのみ押すと終了）:")
    notes = []
    while True:
        note = input(f"  注意事項{len(notes) + 1}: ").strip()
        if not note:
            break
        notes.append(note)
    build_print_layout(master, students, print_out_path, school_name, is_checked_fn,
                        sale_date=sale_date, notes=notes)


def main():
    print("=== 室蘭栄高校 教科書購入票 自動生成 ===")
    print("1: 新2年生")
    print("2: 新3年生")
    print("3: 新1年生（共通リスト）")
    grade = input("学年を選んでください（1/2/3）: ").strip()

    if grade == '1':
        purchase_path = input("購入票ファイル（前年版など）のパスを入力: ").strip()
        raw_path = input("2年生の生データファイルのパスを入力: ").strip()
        out_path = input("出力ファイル名（例: 栄高校_新2年_購入票.xlsx）: ").strip()
        master = load_master_2nen(purchase_path)
        students = load_students_2nen(raw_path)
        build_workbook(master, students, out_path, is_checked_2nen)
        ask_and_build_print_layout(master, students, is_checked_2nen, '室蘭栄高等学校２学年')
    elif grade == '2':
        purchase_path = input("購入票ファイル（前年版など）のパスを入力: ").strip()
        raw_path = input("3年生の生データファイルのパスを入力: ").strip()
        out_path = input("出力ファイル名（例: 栄高校_新3年_購入票.xlsx）: ").strip()
        master = load_master_3nen(purchase_path)
        students = load_students_3nen(raw_path)
        build_workbook(master, students, out_path, is_checked_3nen)
        ask_and_build_print_layout(master, students, is_checked_3nen, '室蘭栄高等学校３学年')
    elif grade == '3':
        purchase_path = input("購入票ファイル（前年版など）のパスを入力: ").strip()
        out_path = input("出力ファイル名（例: 栄高校_新1年_購入リスト.xlsx）: ").strip()
        master, geijutsu = load_1nen(purchase_path)
        build_1nen_workbook(master, geijutsu, out_path)
        opts = printlayout.ask_print_options('栄高校_新1年_購入リスト（印刷用）.xlsx')
        if opts:
            build_1nen_print_layout(master, geijutsu, opts['path'],
                                    sale_date=opts['sale_date'], notes=opts['notes'])
    else:
        print("1・2・3のいずれかを入力してください。")


if __name__ == '__main__':
    main()
